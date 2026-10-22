##
 # @file test/excel_contract/cases.py
 # @brief Deterministic inputs and shared legacy-result serialization.
 #

import datetime
import math
import os
from pathlib import Path


##
 # @brief Build a workbook without committing binary fixtures.
 #
def workbook_builder(rows, sheet_name="Sheet1", trailing_cells=()):
    def build(directory):
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_name
        for row in rows:
            sheet.append(row)
        # End-for
        for address in trailing_cells:
            # Styled blanks exercise dimensions, not just omitted XML cells.
            sheet[address].number_format = "0.00"
        # End-for
        workbook.save(Path(directory) / "input.xlsx")
        workbook.close()
    # End-def
    return build
# End-def


##
 # @brief Include a second sheet for index/name selection.
 #
def multiple_sheets(directory):
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.active.append(["first"])
    workbook.active.append([10])
    sheet = workbook.create_sheet("Target")
    sheet.append(["second", "value"])
    sheet.append(["selected", 20])
    workbook.save(Path(directory) / "input.xlsx")
    workbook.close()
# End-def


##
 # @brief Explicit Excel errors; openpyxl marks these cells as type 'e'.
 #
def error_cells(directory):
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["error", "number", "label"])
    sheet.append(["#DIV/0!", 1, "first"])
    sheet.append(["#N/A", 2, "second"])
    sheet["A2"].data_type = "e"
    sheet["A3"].data_type = "e"
    workbook.save(Path(directory) / "input.xlsx")
    workbook.close()
# End-def


##
 # @brief Write literal strings even when Excel recognizes an error spelling.
 #
def na_strings(directory):
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["value", "row"])
    for index, text in enumerate(NA_STRINGS):
        sheet.append([text, index + 1])
        sheet.cell(index + 2, 1).data_type = "s"
    # End-for
    workbook.save(Path(directory) / "input.xlsx")
    workbook.close()
# End-def


##
 # @brief Non-workbook files keep error text independent of ZIP internals.
 #
def text_files(directory):
    for name in ("input.txt", "input.xls"):
        (Path(directory) / name).write_text("not an Excel workbook\n", encoding="ascii")
    # End-for
# End-def


##
 # @brief Prove successful writes replace, rather than append to, workbooks.
 #
def existing_output(directory):
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.active.title = "OldSheet"
    workbook.active.append(["old", "content"])
    workbook.create_sheet("AlsoOld")
    workbook.save(Path(directory) / "output.xlsx")
    workbook.close()
# End-def


NA_STRINGS = [
    "NA", "N/A", "NULL", "nan", "None", "#N/A", "n/a", "null", "-nan", "<NA>",
    "", "#NA", "#N/A N/A", "-1.#IND", "-1.#QNAN", "1.#IND", "1.#QNAN", "NaN", "-NaN",
]
BASIC_ROWS = [["name", "count"], ["apple", 1], ["pear", 2]]
BASIC_TABLE = "| name | count |\n| --- | --- |\n| apple | 1 |\n| pear | 2 |"
CASES = []


##
 # @brief Add a read contract with its own independently rebuilt input.
 #
def read_case(case_id, rows=BASIC_ROWS, builder=None, **kwargs):
    CASES.append({
        "id": "read_" + case_id,
        "builder": builder if builder is not None else workbook_builder(rows),
        "tool": "read_excel",
        "kwargs": {"file_path": "input.xlsx", **kwargs},
    })
# End-def


##
 # @brief Add a write contract, including parsing failures and overwrite cases.
 #
def write_case(case_id, table=BASIC_TABLE, builder=None, **kwargs):
    CASES.append({
        "id": "write_" + case_id,
        "builder": builder,
        "tool": "write_excel",
        "kwargs": {"file_path": "output.xlsx", "markdown_table": table, **kwargs},
    })
# End-def


read_case("basic")
read_case("empty", [])
read_case("header_only", [["name", "count"]])
read_case("sparse", [["a", "b", "c"], [1, None, "x"], [None, 2, None], [3, None, "z"]])
read_case("leading_empty", [[None, None], ["a", "b"], [1, 2]])
read_case("trailing_empty", builder=workbook_builder(BASIC_ROWS, trailing_cells=("G2", "B8", "H10")))
read_case("middle_empty", [["a", "b"], [1, 2], [None, None], [3, 4]])
read_case("dates", [["day"], [datetime.datetime(2024, 2, 3)], [datetime.date(2024, 2, 4)]])
read_case("datetimes", [["when"], [datetime.datetime(2024, 2, 3, 4, 5, 6, 123456)], [datetime.datetime(2024, 2, 4, 12, 30)]])
read_case("times", [["time"], [datetime.time(4, 5, 6)], [datetime.time(12, 30, 1, 123456)]])
read_case("datetime_missing", [["when", "row"], [datetime.datetime(2024, 2, 3), 1], [None, 2]])
read_case("datetime_mixed", [["when"], [datetime.datetime(2024, 2, 3)], ["plain"]])
read_case("booleans", [["flag"], [True], [False]])
read_case("booleans_missing", [["flag", "row"], [True, 1], [None, 2], [False, 3]])
read_case("booleans_numbers", [["flag", "number"], [True, 1], [False, 2]])
read_case("boolean_mixed", [["value"], [True], [2], [False]])
read_case("integers", [["number"], [1], [2], [-3], [12345678901234]])
read_case("floats", [["number"], [1.5], [-2.25], [2.0], [0.000123456789]])
read_case("integral_floats", [["number"], [2.0], [-3.0]])
read_case("numeric_columns", [["int", "float"], [1, 2.5], [12345678901234, -1.25]])
read_case("numeric_missing", [["number", "label"], [1, "a"], [None, "b"], [3, "c"]])
read_case("numeric_missing_only", [["a", "b"], [1, 2], [None, 4], [3, 5]])
read_case("float_missing", [["number", "label"], [1.0, "a"], [None, "b"], [2.5, "c"]])
read_case("numeric_strings", [["zeros", "exponent", "boolean"], ["001", "1e5", "TRUE"], ["002", "2e5", "FALSE"]])
read_case("string_zeros", [["value"], ["001"], ["002"]])
read_case("string_numbers_na", [["value", "row"], ["001", 1], ["NA", 2], ["002", 3]])
read_case("string_float", [["value"], ["1.5"], ["2.0"]])
read_case("string_infinity", [["value"], ["inf"], ["-inf"]])
read_case("numeric_strings_mixed", [["value"], ["001"], ["text"], ["1e5"], ["TRUE"]])
read_case("na_data", builder=na_strings)
read_case("na_headers", [NA_STRINGS, list(range(len(NA_STRINGS)))])
read_case("na_whitespace", [["value"], [" NA "], ["NAN"], ["none"], ["None"], ["ok"]])
read_case("error_cells", builder=error_cells)
read_case("mixed_column", [["mixed"], ["word"], [1], [1.5], [False]])
read_case("duplicate_headers", [["a", "a", "a"], [1, 2, 3]])
read_case("duplicate_collisions", [["a", "a", "a.1", "a", "a.2"], [1, 2, 3, 4, 5]])
read_case("none_headers", [[None, "b", None], [1, 2, 3]])
read_case("all_none_headers", [[None, None], [1, 2]])
read_case("unnamed_collision", [[None, "Unnamed: 0", "a", "a"], [1, 2, 3, 4]])
read_case("numeric_headers", [[1, 2.5, 2.0, True], ["a", "b", "c", "d"]])
read_case("datetime_header", [[datetime.datetime(2024, 2, 3), "value"], ["date header", 1]])
read_case("unicode", [["\u540d\u524d", "value"], ["\u4f60\u597d", "\u6771\u4eac"], ["\ud55c\uae00", "\u6d4b\u8bd5"]])
read_case("multiline_pipe", [["text", "value"], ["first\nsecond | part", 1]])
read_case("sheet_name", builder=multiple_sheets, sheet="Target")
read_case("sheet_index", builder=multiple_sheets, sheet=1)
read_case("sheet_missing", sheet="Missing")
read_case("sheet_out_of_range", sheet=5)
read_case("sheet_negative_index", builder=multiple_sheets, sheet=-1)
read_case("sheet_none", sheet=None)
read_case("header_none", header_row=None)
read_case("header_two", [["skip"], ["also skip"], ["a", "b"], [1, 2], [3, 4]], header_row=2)
read_case("header_out_of_range", header_row=50)
read_case("header_negative", header_row=-1)
read_case("max_none", max_rows=None)
read_case("max_zero", max_rows=0)
read_case("max_two", [["a"], [1], [2], [3], [4]], max_rows=2)
read_case("max_negative", max_rows=-1)
read_case("max_middle_empty", [["a"], [1], [None], [3], [4]], max_rows=2)
read_case("max_type_error", max_rows="2")
read_case("header_type_error", header_row="0")
read_case("missing_file", file_path="missing.xlsx")
read_case("text_file", builder=text_files, file_path="input.txt")
read_case("fake_xls", builder=text_files, file_path="input.xls")
CASES.append({"id": "read_missing_path", "builder": None, "tool": "read_excel", "kwargs": {}})

write_case("basic")
write_case("alignment", "| a | b | c |\n| :--- | ---: | :---: |\n| x | y | z |")
write_case("string_values", "| zero | big | bool | float | empty |\n| --- | --- | --- | --- | --- |\n| 001 | 12345678901234567890 | TRUE | 1.5 | |")
write_case("header_only", "| a | b |\n| --- | --- |")
write_case("short_rows", "| a | b | c |\n| --- | --- | --- |\n| 1 | 2 | 3 |\n| short |")
write_case("all_short_rows", "| a | b | c |\n| --- | --- | --- |\n| short |")
write_case("long_rows", "| a | b |\n| --- | --- |\n| 1 | 2 | 3 |")
write_case("empty", "")
write_case("whitespace_only", "  \n \n ")
write_case("alignment_only", "| --- | --- |")
write_case("missing_path", file_path="")
write_case("custom_sheet", sheet_name="Report")
write_case("overwrite", builder=existing_output)
write_case("nested", file_path="nested/more/output.xlsx")
write_case("whitespace", "  |  a  | b  |  \n| --- | --- |\n|   hello  |  world  |\n\n | 001 |   | ")
write_case("without_pipes", "a | b\n--- | ---\nx | y")
write_case("without_pipes_no_rule", "a | b\nx | y")
write_case("one_sided_pipes", "| a | b\n| x | y")
write_case("duplicate_headers", "| a | a |\n| --- | --- |\n| 1 | 2 |")
write_case("empty_header", "| | b |\n| --- | --- |\n| x | y |")
write_case("unicode", "| \u540d\u524d | value |\n| --- | --- |\n| \u4f60\u597d | \u6771\u4eac |")
write_case("formula_error_strings", "| formula | error |\n| --- | --- |\n| =1+1 | #N/A |")
write_case("invalid_sheet", sheet_name="bad/name")


##
 # @brief Replace only the sandbox root embedded in user-visible messages.
 #
def normalize_text(text, directory):
    return text.replace(os.path.abspath(directory), "<TMP>")
# End-def


##
 # @brief JSON-safe, explicitly typed values distinguish text from numbers.
 #
def typed_value(value):
    name = type(value).__name__
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        value = value.isoformat()
    elif isinstance(value, float) and not math.isfinite(value):
        value = str(value)
    # End-if
    return [name, value]
# End-def


##
 # @brief Snapshot saved values and the pandas header cell formatting.
 #
def snapshot_workbook(path):
    from openpyxl import load_workbook
    workbook = load_workbook(path, data_only=False)
    try:
        sheets = []
        for sheet in workbook:
            header_style = []
            for cell in sheet[1]:
                header_style.append({
                    "bold": cell.font.bold,
                    "borders": {
                        side: getattr(cell.border, side).style
                        for side in ("left", "right", "top", "bottom")
                    },
                    "horizontal": cell.alignment.horizontal,
                })
            # End-for
            sheets.append({
                "name": sheet.title,
                "values": [[typed_value(cell.value) for cell in row] for row in sheet],
                "header_style": header_style,
            })
        # End-for
        return {"sheet_names": workbook.sheetnames, "sheets": sheets}
    finally:
        workbook.close()
    # End-try
# End-def


##
 # @brief Invoke either implementation, also characterizing uncaught exceptions.
 #
def run_case(case, directory, tools):
    if case["builder"] is not None:
        case["builder"](directory)
    # End-if
    tool = tools[case["tool"]](workspace_dir=str(directory))
    try:
        success, text = tool.execute(**case["kwargs"])
        result = {"success": success, "text": normalize_text(text, directory)}
    except Exception as error:
        result = {"raises": normalize_text(f"{type(error).__name__}: {error}", directory)}
    # End-try
    if case["tool"] == "write_excel" and result.get("success"):
        path = Path(directory) / case["kwargs"]["file_path"]
        result["workbook"] = snapshot_workbook(path)
    # End-if
    return result
# End-def
