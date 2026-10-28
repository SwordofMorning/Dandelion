##
 # @file src/core/tool/editor/excel_tool.py
 # @date 2026/08/13
 # 
 # @brief Excel Tool, with 2 class.
 # A general excel (xlsx) read/write tool powered by openpyxl and tabulate.
 #

import importlib
import os
import re
from ..base_tool import BaseTool


##
 # @brief Load openpyxl and tabulate with cache-invalidation retry.
 #
def _load_libraries():
    try:
        import openpyxl
        from tabulate import tabulate
    except ImportError:
        importlib.invalidate_caches()
        import openpyxl
        from tabulate import tabulate
    # End-try
    return openpyxl, tabulate
# End-def


##
 # @brief Normalize a single openpyxl cell value for tabular display.
 #
def _cell_value(cell):
    val = cell.value
    if val is None or getattr(cell, "data_type", None) == "e":
        return ""
    # End-if
    if isinstance(val, float) and val.is_integer():
        return int(val)
    # End-if
    return val
# End-def


##
 # @brief Read non-empty rows from an openpyxl worksheet.
 #
def _read_sheet_rows(sheet, header_row, max_rows):
    sheet.reset_dimensions()
    limit = None if max_rows is None else (0 if header_row is None else header_row + 1) + max_rows
    rows = []
    last_non_empty = 0
    for cells in sheet.rows:
        row = [_cell_value(c) for c in cells]
        while row and row[-1] == "":
            row.pop()
        # End-while
        rows.append(row)
        if row:
            last_non_empty = len(rows)
        # End-if
        if limit is not None and len(rows) >= limit:
            break
        # End-if
    # End-for
    rows = rows[:last_non_empty]
    width = max((len(r) for r in rows), default=0)
    return [r + [""] * (width - len(r)) for r in rows]
# End-def


##
 # @brief Excel Read Class.
 #
class ReadExcelTool(BaseTool):
    ##
     # @brief Constructor.
     #
     # @param workspace_dir Default to current directory if not explicitly provided.
     #
    def __init__(self, workspace_dir=None):
        super().__init__(workspace_dir)
    # End-def

    ##
     # @brief Return tool's name.
     #
    def get_name(self):
        return "read_excel"
    # End-def

    ##
     # @brief Return tool's description.
     #
    def get_description(self):
        return (
            "Read an Excel file (.xlsx) and return the specified worksheet as a Markdown table. "
            "Legacy .xls files are not supported. Use this for general tabular data extraction. "
            "Parameters: file_path (required), sheet (name or index, default 0), "
            "header_row (0-based index or null, default 0), "
            "max_rows (limit data rows to save tokens, default 100)."
        )
    # End-def

    ##
     # @brief Return tool's schema.
     #
    def get_schema(self):
        return {
            "type": "object",
            "properties": {
                "file_path": {"type": "string", "description": "Absolute or relative path to the Excel file."},
                "sheet": {"type": ["string", "integer", "null"], "description": "Sheet name or 0-based index. Default: 0."},
                "header_row": {"type": ["integer", "null"], "description": "Row index for headers. Null = no header. Default: 0."},
                "max_rows": {"type": ["integer", "null"], "description": "Max data rows to read. Default: 100."}
            },
            "required": ["file_path"]
        }
    # End-def

    ##
     # @brief Run tool.
     #
     # @param kwargs schema properties.
     #
     # @return (success_bool, result_string)
     #
    def execute(self, **kwargs):
        try:
            openpyxl, tabulate = _load_libraries()
        except ImportError as ex:
            return False, f"Error: Missing libraries ({str(ex)}). Run: pip install openpyxl tabulate"
        # End-try

        file_path = kwargs.get("file_path", "")
        if not file_path:
            return False, "Error: No file_path provided."

        if not os.path.isabs(file_path):
            file_path = os.path.join(self.workspace_dir, file_path)
        file_path = os.path.abspath(file_path)

        # SECURITY: interactive approval + fail-safe re-verify on resolved path.
        resolved, err = self._prepare_path(file_path, action_desc="READ Excel File")
        if err:
            return False, err
        # End-if

        if not os.path.exists(resolved):
            return False, f"Error: File not found: {file_path}"

        sheet = kwargs.get("sheet", 0)
        if sheet is None:
            sheet = 0
        # End-if
        header_row = kwargs.get("header_row", 0)
        max_rows = kwargs.get("max_rows", 100)

        wb = None
        try:
            wb = openpyxl.load_workbook(resolved, read_only=True, data_only=True)
            if isinstance(sheet, str):
                if sheet not in wb.sheetnames:
                    return False, f"Error reading Excel: Worksheet named '{sheet}' not found."
                # End-if
                ws = wb[sheet]
            else:
                if sheet < 0 or sheet >= len(wb.worksheets):
                    return False, f"Error reading Excel: Worksheet index {sheet} out of range."
                # End-if
                ws = wb.worksheets[sheet]
            # End-if

            all_rows = _read_sheet_rows(ws, header_row, max_rows)
        except Exception as e:
            return False, f"Error reading Excel: {str(e)}"
        finally:
            if wb is not None:
                wb.close()
            # End-if
        # End-try

        if not all_rows:
            return True, "(Empty table: no data rows found)"
        # End-if

        if header_row is not None:
            if header_row < 0 or header_row >= len(all_rows):
                return True, "(Empty table: no data rows found)"
            # End-if
            headers = [str(h) if h != "" else f"col_{i}" for i, h in enumerate(all_rows[header_row])]
            data_rows = all_rows[header_row + 1:]
        else:
            headers = ()
            data_rows = all_rows
        # End-if

        if max_rows is not None:
            data_rows = data_rows[:max_rows]
        # End-if

        if not data_rows or not data_rows[0]:
            return True, "(Empty table: no data rows found)"
        # End-if

        try:
            table_str = tabulate(
                data_rows,
                headers=headers,
                tablefmt="github",
                showindex=False
            )
        except Exception as e:
            return False, f"Error formatting table: {str(e)}"
        # End-try

        meta = [
            "[Excel Table Extract]",
            f"File: {os.path.basename(file_path)}",
            f"Shape: {len(data_rows)} rows x {len(data_rows[0])} cols",
            ""
        ]
        return True, "\n".join(meta) + "\n" + table_str
    # End-def execute
# End-class


##
 # @brief Excel Write Class.
 #
class WriteExcelTool(BaseTool):
    ##
     # @brief Constructor.
     #
     # @param workspace_dir Default to current directory if not explicitly provided.
     #
    def __init__(self, workspace_dir=None):
        super().__init__(workspace_dir)
    # End-def

    ##
     # @brief Return tool's name.
     #
    def get_name(self):
        return "write_excel"
    # End-def

    ##
     # @brief Return tool's description.
     #
    def get_description(self):
        return (
            "Write tabular data to an Excel file (.xlsx) from a Markdown table string. "
            "Overwrites the file if it exists."
        )
    # End-def

    ##
     # @brief Return tool's schema.
     #
    def get_schema(self):
        return {
            "type": "object",
            "properties": {
                "file_path": {"type": "string", "description": "Absolute or relative path to output Excel file."},
                "markdown_table": {"type": "string", "description": "Markdown table string (GitHub flavored)."},
                "sheet_name": {"type": "string", "description": "Worksheet name. Default: 'Sheet1'."}
            },
            "required": ["file_path", "markdown_table"]
        }
    # End-def

    ##
     # @brief Run tool.
     #
     # @param kwargs schema properties.
     #
     # @return (success_bool, result_string)
     #
    def execute(self, **kwargs):
        try:
            openpyxl, _ = _load_libraries()
        except ImportError as ex:
            return False, f"Error: Missing libraries ({str(ex)}). Run: pip install openpyxl"
        # End-try

        file_path = kwargs.get("file_path", "")
        md_table = kwargs.get("markdown_table", "")
        sheet_name = kwargs.get("sheet_name", "Sheet1")

        if not file_path or not md_table:
            return False, "Error: file_path and markdown_table are required."

        if not os.path.isabs(file_path):
            file_path = os.path.join(self.workspace_dir, file_path)
        file_path = os.path.abspath(file_path)

        # SECURITY: interactive approval + fail-safe re-verify on resolved path.
        resolved, err = self._prepare_path(file_path, action_desc="WRITE Excel File")
        if err:
            return False, err
        # End-if

        wb = None
        try:
            headers, rows = self._parse_markdown_table(md_table)
            if not headers or not rows:
                return False, "Error: No valid data parsed from markdown table."
            # End-if

            os.makedirs(os.path.dirname(resolved), exist_ok=True)
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = sheet_name
            ws.append(headers)
            for row in rows:
                ws.append(row)
            # End-for
            wb.save(resolved)
            return True, f"Successfully wrote {len(rows)} rows to '{file_path}'."
        except Exception as e:
            return False, f"Error writing Excel: {str(e)}"
        finally:
            if wb is not None:
                wb.close()
            # End-if
        # End-try
    # End-def execute

    ##
     # @brief Parse a GitHub-flavored Markdown table into (headers, rows).
     #
     # @param md Markdown table string (GFM).
     # @return (headers_list, rows_list)
     #
    def _parse_markdown_table(self, md: str):
        lines = [line.strip() for line in md.strip().split("\n") if line.strip()]
        data_lines = [line for line in lines if not re.match(r"^\|[\s\-\|:]+\|$", line)]
        if len(data_lines) < 2:
            return [], []
        # End-if

        parsed = []
        for line in data_lines:
            if line.startswith("|"):
                line = line[1:]
            # End-if
            if line.endswith("|"):
                line = line[:-1]
            # End-if
            parsed.append([cell.strip() for cell in line.split("|")])
        # End-for

        headers = parsed[0]
        width = len(headers)
        rows = []
        for row in parsed[1:]:
            if len(row) < width:
                row = row + [""] * (width - len(row))
            elif len(row) > width:
                row = row[:width]
            # End-if
            rows.append(row)
        # End-for
        return headers, rows
    # End-def
# End-class
