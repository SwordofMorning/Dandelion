##
 # @file src/tool/editor/excel_tool.py
 # @date 2026/08/13
 # 
 # @brief Excel Tool, with 2 class.
 # Read .xlsx files and write Markdown tables with openpyxl and tabulate.
 #

import datetime
import importlib
import math
import os
import re
import zipfile
from ..base_tool import BaseTool

_NA_STRINGS = frozenset((
    "", "#N/A", "#N/A N/A", "#NA", "-1.#IND", "-1.#QNAN", "-NaN", "-nan",
    "1.#IND", "1.#QNAN", "<NA>", "N/A", "NA", "NULL", "NaN", "None", "n/a", "nan", "null",
))
_NUMBER = re.compile(r"^[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?$")


##
 # @brief Load only the required libraries (openpyxl + tabulate).
 #
 # Retains the legacy cache-invalidation retry for packages installed while
 # the agent is running. pandas/numpy are NOT required: openpyxl only probes
 # numpy as an optional extra, and the release artifact does not ship it.
 #
 # @return (openpyxl module, tabulate function); raises ImportError.
 #
def _load_libraries():
    try:
        import openpyxl
        from tabulate import tabulate
    except ImportError:
        importlib.invalidate_caches()
        import openpyxl
        from tabulate import tabulate
    # End-try
    return openpyxl, tabulate
# End-def


##
 # @brief Reproduce the legacy cell reader before column inference.
 #
def _cell_value(cell):
    value = cell.value
    if value is None:
        return ""
    elif cell.data_type == "e":
        return float("nan")
    elif cell.data_type == "n":
        integer = int(value)
        return integer if integer == value else float(value)
    # End-if
    return value
# End-def


##
 # @brief Trim blanks exactly as the legacy openpyxl reader did, even at nrows.
 #
def _sheet_rows(sheet, header_row, max_rows):
    sheet.reset_dimensions()
    limit = None if max_rows is None else (1 if header_row is None else header_row + 1) + max_rows
    rows = []
    last_data = 0
    for cells in sheet.rows:
        row = [_cell_value(cell) for cell in cells]
        while row and row[-1] == "":
            row.pop()
        # End-while
        rows.append(row)
        if row:
            last_data = len(rows)
        # End-if
        if limit is not None and len(rows) >= limit:
            break
        # End-if
    # End-for
    rows = rows[:last_data]
    width = max((len(row) for row in rows), default=0)
    return [row + [""] * (width - len(row)) for row in rows]
# End-def


##
 # @brief Reserve explicit names before filling/mangling unnamed and duplicate headers.
 #
def _column_names(row):
    unnamed = [index for index, value in enumerate(row) if value == ""]
    names = [f"Unnamed: {index}" if value == "" else value for index, value in enumerate(row)]
    counts = {}
    order = [index for index in range(len(names)) if index not in unnamed] + unnamed
    for index in order:
        name = names[index]
        original = name
        count = counts.get(name, 0)
        while count > 0:
            counts[original] = count + 1
            name = f"{original}.{count}"
            count = count + 1 if name in names else counts.get(name, 0)
        # End-while
        names[index] = name
        counts[name] = count + 1
    # End-for
    return [str(name) for name in names]
# End-def


##
 # @brief Preserve the legacy NA vocabulary without stripping text first.
 #
def _is_missing(value):
    return value is None or (isinstance(value, str) and value in _NA_STRINGS) or (
        isinstance(value, float) and math.isnan(value)
    )
# End-def


##
 # @brief Infer each column, including missing-value promotion and boolean strings.
 #
def _infer_column(values):
    values = [None if _is_missing(value) else value for value in values]
    present = [value for value in values if value is not None]
    missing = len(present) != len(values)
    if present and all(isinstance(value, datetime.datetime) for value in present):
        return values, "datetime"
    # End-if
    if present and not missing and all(type(value) is bool for value in present):
        return values, "bool"
    # End-if
    numeric = []
    for value in values:
        if value is None or isinstance(value, (int, float)):
            numeric.append(value)
        elif isinstance(value, str) and _NUMBER.fullmatch(value.strip()):
            text = value.strip()
            numeric.append(float(text) if any(char in text for char in ".eE") else int(text))
        elif isinstance(value, str) and value.strip().lower() in ("inf", "+inf", "-inf", "infinity", "+infinity", "-infinity"):
            numeric.append(float(value))
        else:
            break
        # End-if
    else:
        floating = missing or any(type(value) is float for value in numeric)
        convert = float if floating else int
        return ["" if value is None else convert(value) for value in numeric], "object" if missing else ("float" if floating else "int")
    # End-for
    booleans = {"True": True, "TRUE": True, "true": True, "False": False, "FALSE": False, "false": False}
    if present and all(type(value) is bool or (isinstance(value, str) and value in booleans) for value in present):
        return ["" if value is None else booleans.get(value, value) for value in values], "object" if missing else "bool"
    # End-if
    return ["" if value is None else value for value in values], "object"
# End-def


##
 # @brief Match the scalar promotion of the legacy DataFrame.values matrix.
 #
def _table_rows(rows):
    columns = [_infer_column(list(values)) for values in zip(*rows)]
    kinds = {kind for values, kind in columns}
    numeric_matrix = kinds <= {"int", "float"}
    date_matrix = kinds == {"datetime"}
    bool_matrix = kinds == {"bool"}
    converted = []
    for values, kind in columns:
        if kind == "datetime":
            values = [
                ("NaT" if value is None else value.isoformat(timespec="microseconds"))
                if date_matrix else ("NaT" if value is None else value)
                for value in values
            ]
        elif (numeric_matrix and "float" in kinds) or bool_matrix:
            # numpy bool scalars are numeric to tabulate; Python bools are not.
            values = [float(value) for value in values]
        # End-if
        converted.append(values)
    # End-for
    return [list(row) for row in zip(*converted)]
# End-def

##
 # @brief Excel Read Class.
 #
class ReadExcelTool(BaseTool):
    ##
     # @brief Constructor.
     #
     # @param workspace_dir Default to current directory if not explicitly provided.
     #
    def __init__(self, workspace_dir=None):
        super().__init__(workspace_dir)
    # End-def

    ##
     # @brief Return tool's name.
     #
    def get_name(self):
        return "read_excel"
    # End-def

    ##
     # @brief Return tool's description.
     #
    def get_description(self):
        return (
            "Read an Excel file (.xlsx) and return the specified worksheet as a Markdown table. "
            "Legacy .xls files are not supported. Use this for general tabular data extraction. "
            "Parameters: file_path (required), sheet (name or index, default 0), "
            "header_row (0-based index or null, default 0), "
            "max_rows (limit data rows to save tokens, default 100)."
        )
    # End-def

    ##
     # @brief Return tool's schema.
     #
    def get_schema(self):
        return {
            "type": "object",
            "properties": {
                "file_path": {"type": "string", "description": "Absolute or relative path to the Excel file."},
                "sheet": {"type": ["string", "integer", "null"], "description": "Sheet name or 0-based index. Default: 0."},
                "header_row": {"type": ["integer", "null"], "description": "Row index for headers. Null = no header. Default: 0."},
                "max_rows": {"type": ["integer", "null"], "description": "Max data rows to read. Default: 100."}
            },
            "required": ["file_path"]
        }
    # End-def

    ##
     # @brief Run tool.
     #
     # @param kwargs schema properties.
     #
     # @return (success_bool, result_string)
     #
    def execute(self, **kwargs):
        # Dynamic Import
        try:
            openpyxl, tabulate = _load_libraries()
        except ImportError as error:
            return False, f"Error: Missing libraries ({str(error)}). Run: pip install openpyxl tabulate"
        # End-try

        file_path = kwargs.get("file_path", "")
        if not file_path:
            return False, "Error: No file_path provided."

        if not os.path.isabs(file_path):
            file_path = os.path.join(self.workspace_dir, file_path)
        file_path = os.path.abspath(file_path)

        # SECURITY: interactive approval + fail-safe re-verify on resolved path.
        resolved, err = self._prepare_path(file_path, action_desc="READ Excel File")
        if err:
            return False, err
        # End-if

        if not os.path.exists(resolved):
            return False, f"Error: File not found: {file_path}"

        sheet = kwargs.get("sheet", 0)
        header_row = kwargs.get("header_row", 0)
        max_rows = kwargs.get("max_rows", 100)

        if sheet is None:
            return False, "Error: 'sheet' must be a sheet name or 0-based index; reading all sheets is not supported."
        # End-if
        workbook = None
        try:
            if header_row is not None:
                if not isinstance(header_row, int):
                    raise ValueError("header must be integer or list of integers")
                elif header_row < 0:
                    raise ValueError("Passing negative integer to header is invalid. For no header, use header=None instead")
                # End-if
            # End-if
            if max_rows is not None and (not isinstance(max_rows, int) or max_rows < 0):
                raise ValueError("'nrows' must be an integer >=0")
            # End-if
            # pandas supplied an open binary handle, so invalid .txt/.xls inputs
            # reached ZIP validation rather than openpyxl's extension rejection.
            if not zipfile.is_zipfile(resolved):
                raise zipfile.BadZipFile("File is not a zip file")
            # End-if
            workbook = openpyxl.load_workbook(resolved, read_only=True, data_only=True)
            if isinstance(sheet, str):
                if sheet not in workbook.sheetnames:
                    raise ValueError(f"Worksheet named '{sheet}' not found")
                # End-if
                worksheet = workbook[sheet]
            else:
                if sheet >= len(workbook.worksheets):
                    raise ValueError(f"Worksheet index {sheet} is invalid, {len(workbook.worksheets)} worksheets found")
                # End-if
                worksheet = workbook.worksheets[sheet]
            # End-if
            rows = _sheet_rows(worksheet, header_row, max_rows)
            if not rows:
                return True, "(Empty table: no data rows found)"
            # End-if
            if header_row is not None and header_row >= len(rows):
                raise ValueError(f"Passed header=[{header_row}], len of 1, but only {len(rows)} lines in file (sheet: {sheet})")
            # End-if
            headers = _column_names(rows[header_row]) if header_row is not None else ()
            rows = rows[header_row + 1:] if header_row is not None else rows
            if max_rows is not None:
                rows = rows[:max_rows]
            # End-if
        except Exception as error:
            return False, f"Error reading Excel: {str(error)}"
        finally:
            if workbook is not None:
                workbook.close()
            # End-if
        # End-try

        if not rows or not rows[0]:
            return True, "(Empty table: no data rows found)"
        # End-if
        try:
            table_str = tabulate(_table_rows(rows), headers=headers, tablefmt="github", showindex=False)
        except Exception as error:
            return False, f"Error formatting table: {str(error)}"
        # End-try
        meta = [
            "[Excel Table Extract]",
            f"File: {os.path.basename(file_path)}",
            f"Shape: {len(rows)} rows x {len(rows[0])} cols",
            ""
        ]
        return True, "\n".join(meta) + "\n" + table_str
    # End-def execute
# End-class

##
 # @brief Excel Write Class.
 #
class WriteExcelTool(BaseTool):
    ##
     # @brief Constructor.
     #
     # @param workspace_dir Default to current directory if not explicitly provided.
     #
    def __init__(self, workspace_dir=None):
        super().__init__(workspace_dir)
    # End-def

    ##
     # @brief Return tool's name.
     #
    def get_name(self):
        return "write_excel"
    # End-def

    ##
     # @brief Return tool's description.
     #
    def get_description(self):
        return (
            "Write tabular data to an Excel file (.xlsx) from a Markdown table string. "
            "Overwrites the file if it exists."
        )
    # End-def

    ##
     # @brief Return tool's schema.
     #
    def get_schema(self):
        return {
            "type": "object",
            "properties": {
                "file_path": {"type": "string", "description": "Absolute or relative path to output Excel file."},
                "markdown_table": {"type": "string", "description": "Markdown table string (GitHub flavored)."},
                "sheet_name": {"type": "string", "description": "Worksheet name. Default: 'Sheet1'."}
            },
            "required": ["file_path", "markdown_table"]
        }
    # End-def

    ##
     # @brief Run tool.
     #
     # @param kwargs schema properties.
     #
     # @return (success_bool, result_string)
     #
    def execute(self, **kwargs):
        # Dynamic Import
        try:
            openpyxl, _ = _load_libraries()
        except ImportError as error:
            return False, f"Error: Missing libraries ({str(error)}). Run: pip install openpyxl tabulate"
        # End-try

        file_path = kwargs.get("file_path", "")
        md_table = kwargs.get("markdown_table", "")
        sheet_name = kwargs.get("sheet_name", "Sheet1")

        if not file_path or not md_table:
            return False, "Error: file_path and markdown_table are required."

        if not os.path.isabs(file_path):
            file_path = os.path.join(self.workspace_dir, file_path)
        file_path = os.path.abspath(file_path)

        # SECURITY: interactive approval + fail-safe re-verify on resolved path.
        resolved, err = self._prepare_path(file_path, action_desc="WRITE Excel File")
        if err:
            return False, err
        # End-if

        workbook = None
        try:
            headers, rows = self._parse_markdown_table(md_table)
            if not headers or not rows:
                return False, "Error: No valid data parsed from markdown table."
            # End-if
            os.makedirs(os.path.dirname(resolved), exist_ok=True)
            workbook = openpyxl.Workbook()
            worksheet = workbook.active
            worksheet.title = sheet_name
            worksheet.append(headers)
            for row in rows:
                worksheet.append(row)
            # End-for
            # The pinned pandas version writes plain, unstyled header cells.
            workbook.save(resolved)
            return True, f"Successfully wrote {len(rows)} rows to '{file_path}'."
        except Exception as error:
            return False, f"Error writing Excel: {str(error)}"
        finally:
            if workbook is not None:
                workbook.close()
            # End-if
        # End-try
    # End-def execute

    ##
     # @brief Parse a Markdown table using the legacy stripping and alignment rules.
     #
     # @param md Markdown table string (GFM).
     # @return (headers, rows), preserving text without numeric coercion.
     #
    def _parse_markdown_table(self, md: str):
        lines = [line.strip() for line in md.strip().split("\n") if line.strip()]
        data_lines = [line for line in lines if not re.match(r"^\|[\s\-\|:]+\|$", line)]
        rows = []
        for line in data_lines:
            if line.startswith("|") and line.endswith("|"):
                line = line[1:-1]
            # End-if
            rows.append([cell.strip() for cell in line.split("|")])
        # End-for
        if not rows:
            return [], []
        # End-if
        headers, data = rows[0], rows[1:]
        width = max((len(row) for row in data), default=len(headers))
        if width != len(headers):
            raise ValueError(f"{len(headers)} columns passed, passed data had {width} columns")
        # End-if
        return headers, [row + [""] * (width - len(row)) for row in data]
    # End-def
# End-class